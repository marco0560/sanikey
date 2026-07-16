"""Document inventory and extraction tests."""

from __future__ import annotations

import shutil
import subprocess
import zipfile
from dataclasses import replace
from pathlib import Path

import pytest

import sanikey.documents as documents_module
from sanikey.config import IngestionConfig, PersonConfig
from sanikey.documents import (
    document_page_count,
    duplicate_document_warnings,
    extract_text,
    find_duplicate_documents,
    scan_document_inventory,
    scan_documents,
)


def _person(tmp_path: Path) -> PersonConfig:
    """Build a synthetic patient config.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary test directory.

    Returns
    -------
    PersonConfig
        Patient configuration.
    """

    return PersonConfig(
        id="patient-a",
        display_name="Patient A",
        source_documents=tmp_path / "documents",
        metadata_directory=tmp_path / "metadata",
        local_build=tmp_path / "generated",
        usb_uuid="1A2B-3C4D",
    )


def _write_xlsx(path: Path, rows: tuple[tuple[str, ...], ...]) -> None:
    """Write a minimal XLSX workbook for extraction tests.

    Parameters
    ----------
    path : pathlib.Path
        Output workbook path.
    rows : tuple[tuple[str, ...], ...]
        Worksheet rows to write as inline strings.

    Returns
    -------
    None
    """

    row_xml = []
    for row_index, row in enumerate(rows, start=1):
        cells = []
        for column_index, value in enumerate(row, start=1):
            column = chr(ord("A") + column_index - 1)
            cells.append(
                f'<c r="{column}{row_index}" t="inlineStr"><is><t>{value}</t></is></c>'
            )
        row_xml.append(f'<row r="{row_index}">{"".join(cells)}</row>')
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr(
            "[Content_Types].xml",
            """<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>
<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
</Types>""",
        )
        archive.writestr(
            "_rels/.rels",
            """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>
</Relationships>""",
        )
        archive.writestr(
            "xl/workbook.xml",
            """<?xml version="1.0" encoding="UTF-8"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
<sheets><sheet name="Vitals" sheetId="1" r:id="rId1"/></sheets>
</workbook>""",
        )
        archive.writestr(
            "xl/_rels/workbook.xml.rels",
            """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>
</Relationships>""",
        )
        archive.writestr(
            "xl/worksheets/sheet1.xml",
            f"""<?xml version="1.0" encoding="UTF-8"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
<sheetData>{"".join(row_xml)}</sheetData>
</worksheet>""",
        )


def test_scan_documents_builds_stable_records(tmp_path: Path) -> None:
    """Verify scanner derives category, date, title, kind, and digest.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.

    Returns
    -------
    None
    """

    person = _person(tmp_path)
    document_dir = person.source_documents / "laboratory"
    document_dir.mkdir(parents=True)
    (document_dir / "20260102 Blood Test.txt").write_text(
        "Synthetic text",
        encoding="utf-8",
    )

    documents = scan_documents(person)

    assert len(documents) == 1
    assert documents[0].patient_id == "patient-a"
    assert documents[0].category == "laboratory"
    assert documents[0].date == "2026-01-02"
    assert documents[0].title == "Blood Test"
    assert documents[0].kind == "text"
    assert len(documents[0].sha256) == 64


def test_scan_documents_hides_service_directory_prefix_from_category(
    tmp_path: Path,
) -> None:
    """Verify service directory markers do not leak into document categories.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.

    Returns
    -------
    None
    """

    person = _person(tmp_path)
    document_dir = person.source_documents / "_Parametri"
    document_dir.mkdir(parents=True)
    (document_dir / "20260102 Pressione.txt").write_text(
        "pressione",
        encoding="utf-8",
    )
    archive_dir = person.source_documents / "_Archivi"
    archive_dir.mkdir()
    (archive_dir / "20260103 Studio.iso").write_bytes(b"iso")

    documents = scan_documents(person)

    documents_by_name = {document.path.name: document for document in documents}
    assert documents_by_name["20260102 Pressione.txt"].category == "Parametri"
    assert documents_by_name["20260103 Studio.iso"].category == "Archivi"


def test_scan_documents_applies_ingestion_exclusions(tmp_path: Path) -> None:
    """Verify configured source exclusions run before inventory creation.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.

    Returns
    -------
    None
    """

    base = _person(tmp_path)
    person = replace(
        base,
        ingestion=IngestionConfig(
            exclude_patterns=("**/Help/**", "*.tmp"),
            include_patterns=("**/Help/keep.txt",),
        ),
    )
    person.source_documents.mkdir(parents=True)
    (person.source_documents / "20260102 Report.txt").write_text(
        "include",
        encoding="utf-8",
    )
    help_dir = person.source_documents / "Help"
    help_dir.mkdir(parents=True)
    (help_dir / "manual.txt").write_text("exclude", encoding="utf-8")
    (help_dir / "keep.txt").write_text("include", encoding="utf-8")
    (person.source_documents / "scratch.tmp").write_text("exclude", encoding="utf-8")

    inventory = scan_document_inventory(person)

    assert [document.path.name for document in inventory] == [
        "20260102 Report.txt",
        "keep.txt",
    ]


def test_scan_documents_classifies_archive_and_office_kinds(
    tmp_path: Path,
) -> None:
    """Verify archive and office document kind classification.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.

    Returns
    -------
    None
    """

    person = _person(tmp_path)
    document_dir = person.source_documents
    document_dir.mkdir(parents=True)
    (document_dir / "20260102 Archive.7z").write_bytes(b"archive")
    (document_dir / "20260103 Legacy.rar").write_bytes(b"rar")
    (document_dir / "20260104 Support.zip").write_bytes(b"zip")
    (document_dir / "20260105 Report.docx").write_bytes(b"docx")
    (document_dir / "20260106 Workbook.xlsx").write_bytes(b"xlsx")

    documents = scan_document_inventory(person)
    kinds = {document.path.name: document.kind for document in documents}

    assert kinds == {
        "20260102 Archive.7z": "archive",
        "20260103 Legacy.rar": "archive",
        "20260104 Support.zip": "archive",
        "20260105 Report.docx": "office",
        "20260106 Workbook.xlsx": "office",
    }


def test_scan_documents_classifies_img_disk_images_as_dicom_support(
    tmp_path: Path,
) -> None:
    """Verify IMG disk images are recognized as DICOM support containers.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.

    Returns
    -------
    None
    """

    person = _person(tmp_path)
    document_dir = person.source_documents
    document_dir.mkdir(parents=True)
    (document_dir / "20260102 Study.img").write_bytes(b"synthetic image")

    documents = scan_documents(person)

    assert documents[0].kind == "dicom_img"


def test_scan_documents_detects_dicom_magic_without_extension(
    tmp_path: Path,
) -> None:
    """Verify extensionless DICOM files are classified by magic bytes.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.

    Returns
    -------
    None
    """

    person = _person(tmp_path)
    document_dir = person.source_documents
    document_dir.mkdir(parents=True)
    (document_dir / "IM000001").write_bytes((b"\0" * 128) + b"DICM" + b"synthetic")

    documents = scan_document_inventory(person)

    assert documents[0].kind == "dicom_file"


def test_scan_documents_classifies_images(tmp_path: Path) -> None:
    """Verify clinical image files are first-class image documents.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.

    Returns
    -------
    None
    """

    person = _person(tmp_path)
    document_dir = person.source_documents
    document_dir.mkdir(parents=True)
    (document_dir / "20260102 Photo.jpg").write_bytes(b"image")
    (document_dir / "20260103 Scan.PNG").write_bytes(b"image")

    documents = scan_document_inventory(person)

    assert [document.kind for document in documents] == ["image", "image"]


def test_duplicate_detection_skips_duplicate_content_with_warning(
    tmp_path: Path,
) -> None:
    """Verify duplicate content is reported and skipped for ingestion.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.

    Returns
    -------
    None
    """

    person = _person(tmp_path)
    document_dir = person.source_documents / "reports"
    document_dir.mkdir(parents=True)
    (document_dir / "20260102 A.txt").write_text("same", encoding="utf-8")
    (document_dir / "20260103 B.txt").write_text("same", encoding="utf-8")

    inventory = scan_document_inventory(person)
    documents = scan_documents(person)
    duplicates = find_duplicate_documents(inventory)
    warnings = duplicate_document_warnings(duplicates)

    assert len(inventory) == 2
    assert len(documents) == 1
    assert documents[0].path.name == "20260102 A.txt"
    assert len(duplicates) == 1
    assert len(next(iter(duplicates.values()))) == 2
    warning_lines = warnings[0].splitlines()
    assert len(warning_lines) == 3
    assert warning_lines[0].startswith(
        "contenuto documento duplicato saltato. I file seguenti sono identici (sha256="
    )
    assert warning_lines[1].endswith("20260102 A.txt")
    assert warning_lines[2].endswith("20260103 B.txt")


def test_extract_text_reads_text_files(tmp_path: Path) -> None:
    """Verify supported text files are extracted directly.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.

    Returns
    -------
    None
    """

    person = _person(tmp_path)
    document_dir = person.source_documents
    document_dir.mkdir(parents=True)
    path = document_dir / "note.txt"
    path.write_text("hello", encoding="utf-8")
    document = scan_documents(person)[0]

    extracted = extract_text(document)

    assert extracted.text == "hello"
    assert extracted.warnings == ()
    assert document_page_count(document) == 1


def test_document_page_count_unknown_for_non_paginated_formats(tmp_path: Path) -> None:
    """Verify page count is unknown for non-paginated extracted formats.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.

    Returns
    -------
    None
    """

    person = _person(tmp_path)
    document_dir = person.source_documents
    document_dir.mkdir(parents=True)
    path = document_dir / "20260102 Workbook.xlsx"
    path.write_bytes(b"not a real workbook")
    document = scan_documents(person)[0]

    assert document_page_count(document) is None


def test_extract_text_skips_source_image_ocr(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Verify source images are catalog-only for text extraction.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.
    monkeypatch : pytest.MonkeyPatch
        Pytest monkeypatch fixture.

    Returns
    -------
    None
    """

    def fail_run(*_args, **_kwargs):
        """Fail if image extraction tries to execute a subprocess.

        Parameters
        ----------
        *_args : object
            Positional arguments ignored by the fake.
        **_kwargs : object
            Keyword arguments ignored by the fake.

        Returns
        -------
        None
            This helper always raises.
        """

        msg = "image OCR must not invoke subprocesses"
        raise AssertionError(msg)

    person = _person(tmp_path)
    document_dir = person.source_documents
    document_dir.mkdir(parents=True)
    path = document_dir / "20260102 Photo.png"
    path.write_bytes(b"image")
    document = scan_documents(person)[0]
    monkeypatch.setattr(documents_module.subprocess, "run", fail_run)

    extracted = extract_text(document)

    assert extracted.text == ""
    assert extracted.warnings == ()


def test_extract_text_dicom_is_catalog_only_without_warning(tmp_path: Path) -> None:
    """Verify DICOM documents do not produce extraction warnings.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.

    Returns
    -------
    None
    """

    person = _person(tmp_path)
    document_dir = person.source_documents
    document_dir.mkdir(parents=True)
    path = document_dir / "image.dcm"
    path.write_bytes((b"\0" * 128) + b"DICM")
    document = scan_documents(person)[0]

    extracted = extract_text(document)

    assert extracted.text == ""
    assert extracted.warnings == ()


def test_extract_text_lists_zip_archive_contents(tmp_path: Path) -> None:
    """Verify ZIP files produce an archive inventory.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.

    Returns
    -------
    None
    """

    person = _person(tmp_path)
    document_dir = person.source_documents
    document_dir.mkdir(parents=True)
    path = document_dir / "20260102 Archive.zip"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("report.txt", "hello")
        archive.writestr("nested/image.dcm", "dicom")
    document = scan_documents(person)[0]

    extracted = extract_text(document)

    assert "contenuto archivio zip:" in extracted.text
    assert "nested/image.dcm" in extracted.text
    assert "report.txt" in extracted.text
    assert extracted.warnings == ()


def test_extract_text_lists_7z_archive_contents(tmp_path: Path) -> None:
    """Verify 7z files produce an archive inventory.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.

    Returns
    -------
    None
    """

    import py7zr

    person = _person(tmp_path)
    document_dir = person.source_documents
    payload = tmp_path / "payload"
    payload.mkdir(parents=True)
    (payload / "report.txt").write_text("hello", encoding="utf-8")
    document_dir.mkdir(parents=True)
    path = document_dir / "20260102 Archive.7z"
    with py7zr.SevenZipFile(path, "w") as archive:
        archive.write(payload / "report.txt", "report.txt")
    document = scan_documents(person)[0]

    extracted = extract_text(document)

    assert "contenuto archivio 7z:" in extracted.text
    assert "report.txt" in extracted.text
    assert extracted.warnings == ()


def test_extract_text_reads_docx_documents(tmp_path: Path) -> None:
    """Verify DOCX text extraction.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.

    Returns
    -------
    None
    """

    import docx

    person = _person(tmp_path)
    document_dir = person.source_documents
    document_dir.mkdir(parents=True)
    path = document_dir / "20260102 Report.docx"
    document_file = docx.Document()
    document_file.add_paragraph("Synthetic DOCX text")
    document_file.save(path)
    document = scan_documents(person)[0]

    extracted = extract_text(document)

    assert "Synthetic DOCX text" in extracted.text
    assert extracted.warnings == ()


def test_extract_text_reads_xlsx_workbooks(tmp_path: Path) -> None:
    """Verify XLSX text extraction.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.

    Returns
    -------
    None
    """

    person = _person(tmp_path)
    document_dir = person.source_documents
    document_dir.mkdir(parents=True)
    path = document_dir / "20260102 Workbook.xlsx"
    _write_xlsx(path, (("Peso", "70 kg"),))
    document = scan_documents(person)[0]

    extracted = extract_text(document)

    assert "[Vitals]" in extracted.text
    assert "Peso\t70 kg" in extracted.text
    assert extracted.warnings == ()


def test_extract_text_captures_xlsx_library_warnings(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Verify XLSX library warnings are returned as extraction warnings.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.
    monkeypatch : pytest.MonkeyPatch
        Pytest monkeypatch fixture.

    Returns
    -------
    None
    """

    import warnings

    import python_calamine

    class FakeWorksheet:
        """Minimal worksheet for XLSX extraction tests.

        Parameters
        ----------
        None

        Returns
        -------
        None
        """

        def to_python(self, *, skip_empty_area: bool):
            """Return synthetic row values.

            Parameters
            ----------
            skip_empty_area : bool
                Whether leading empty cells should be skipped.

            Returns
            -------
            list[list[str]]
                Synthetic rows.
            """

            assert skip_empty_area is False
            return [["Synthetic cell"]]

    class FakeWorkbook:
        """Minimal workbook for XLSX extraction tests.

        Parameters
        ----------
        None

        Returns
        -------
        None
        """

        sheet_names = ["Sheet"]

        def get_sheet_by_name(self, name: str) -> FakeWorksheet:
            """Return a synthetic worksheet by name.

            Parameters
            ----------
            name : str
                Requested worksheet name.

            Returns
            -------
            FakeWorksheet
                Synthetic worksheet.
            """

            assert name == "Sheet"
            return FakeWorksheet()

        def close(self) -> None:
            """Close the fake workbook.

            Parameters
            ----------
            None

            Returns
            -------
            None
            """

    def fake_load_workbook(*_args, **_kwargs):
        """Return a workbook while emitting an XLSX parser warning.

        Parameters
        ----------
        *_args : object
            Positional arguments ignored by the fake.
        **_kwargs : object
            Keyword arguments ignored by the fake.

        Returns
        -------
        FakeWorkbook
            Synthetic workbook.
        """

        warnings.warn("Data Validation extension is not supported", stacklevel=2)
        return FakeWorkbook()

    person = _person(tmp_path)
    document_dir = person.source_documents
    document_dir.mkdir(parents=True)
    path = document_dir / "20260102 Workbook.xlsx"
    path.write_bytes(b"xlsx")
    document = scan_documents(person)[0]
    monkeypatch.setattr(python_calamine, "load_workbook", fake_load_workbook)

    extracted = extract_text(document)

    assert "Synthetic cell" in extracted.text
    assert extracted.warnings == (
        "testo foglio di calcolo estratto; funzionalita' di compatibilita' "
        "cartella non preservata: "
        "Data Validation extension is not supported",
    )


def test_extract_text_reads_odt_documents(tmp_path: Path) -> None:
    """Verify ODT text extraction.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.

    Returns
    -------
    None
    """

    from odf.opendocument import OpenDocumentText
    from odf.text import P

    person = _person(tmp_path)
    document_dir = person.source_documents
    document_dir.mkdir(parents=True)
    path = document_dir / "20260102 Report.odt"
    document_file = OpenDocumentText()
    document_file.text.addElement(P(text="Synthetic ODT text"))
    document_file.save(path)
    document = scan_documents(person)[0]

    extracted = extract_text(document)

    assert "Synthetic ODT text" in extracted.text
    assert extracted.warnings == ()


def test_extract_text_reads_ods_spreadsheets(tmp_path: Path) -> None:
    """Verify ODS text extraction.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.

    Returns
    -------
    None
    """

    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table, TableCell, TableRow
    from odf.text import P

    person = _person(tmp_path)
    document_dir = person.source_documents
    document_dir.mkdir(parents=True)
    path = document_dir / "20260102 Spreadsheet.ods"
    document_file = OpenDocumentSpreadsheet()
    table = Table(name="Sheet1")
    row = TableRow()
    cell = TableCell(valuetype="string")
    cell.addElement(P(text="Synthetic ODS text"))
    row.addElement(cell)
    table.addElement(row)
    document_file.spreadsheet.addElement(table)
    document_file.save(path)
    document = scan_documents(person)[0]

    extracted = extract_text(document)

    assert "Synthetic ODS text" in extracted.text
    assert extracted.warnings == ()


def test_extract_text_reads_legacy_office_through_libreoffice(
    tmp_path: Path,
    monkeypatch,
) -> None:
    """Verify legacy Office extraction uses LibreOffice output.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.
    monkeypatch : pytest.MonkeyPatch
        Pytest monkeypatch fixture.

    Returns
    -------
    None
    """

    person = _person(tmp_path)
    document_dir = person.source_documents
    document_dir.mkdir(parents=True)
    path = document_dir / "20260102 Legacy.doc"
    path.write_bytes(b"legacy")
    document = scan_documents(person)[0]

    def fake_run(command, **_kwargs):
        output_dir = Path(command[command.index("--outdir") + 1])
        (output_dir / "20260102 Legacy.txt").write_text(
            "legacy office text",
            encoding="utf-8",
        )
        return subprocess.CompletedProcess(command, 0, "", "")

    monkeypatch.setattr(documents_module.shutil, "which", lambda _: "/usr/bin/soffice")
    monkeypatch.setattr(documents_module.subprocess, "run", fake_run)

    extracted = extract_text(document)

    assert extracted.text == "legacy office text"
    assert extracted.warnings == ()


@pytest.mark.parametrize("extension", [".docx", ".xlsx", ".xls", ".ods", ".odt"])
def test_extract_text_warns_for_corrupt_office_documents(
    tmp_path: Path,
    extension: str,
) -> None:
    """Verify corrupt office-like files produce warnings.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.
    extension : str
        File extension under test.

    Returns
    -------
    None
    """

    person = _person(tmp_path)
    document_dir = person.source_documents
    document_dir.mkdir(parents=True)
    path = document_dir / f"20260102 Broken{extension}"
    path.write_bytes(b"not a valid office document")
    document = scan_documents(person)[0]

    extracted = extract_text(document)

    assert extracted.text == ""
    assert extracted.warnings


def test_extract_text_skips_dicom_support_without_warning(tmp_path: Path) -> None:
    """Verify DICOM support files are catalog-only without warnings.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.

    Returns
    -------
    None
    """

    person = _person(tmp_path)
    document_dir = person.source_documents
    document_dir.mkdir(parents=True)
    path = document_dir / "20260102 Study.iso"
    path.write_bytes(b"not a real iso")
    document = scan_documents(person)[0]

    extracted = extract_text(replace(document, path=path))

    assert extracted.text == ""
    assert extracted.warnings == ()


def test_extract_text_uses_ocrmypdf_when_pymupdf_is_unavailable(
    tmp_path: Path,
    monkeypatch,
) -> None:
    """Verify system OCRmyPDF is a PDF extraction fallback.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.
    monkeypatch : pytest.MonkeyPatch
        Pytest monkeypatch fixture.

    Returns
    -------
    None
    """

    person = _person(tmp_path)
    document_dir = person.source_documents
    document_dir.mkdir(parents=True)
    path = document_dir / "20260102 Report.pdf"
    path.write_bytes(b"%PDF synthetic")
    document = scan_documents(person)[0]

    def fake_run(command, **_kwargs):
        sidecar = Path(command[command.index("--sidecar") + 1])
        sidecar.write_text("ocr text", encoding="utf-8")
        return subprocess.CompletedProcess(command, 0, "", "")

    monkeypatch.setattr(
        documents_module, "_extract_pdf_text_with_pymupdf", lambda _: None
    )
    monkeypatch.setattr(documents_module.shutil, "which", lambda _: "/usr/bin/ocrmypdf")
    monkeypatch.setattr(documents_module.subprocess, "run", fake_run)

    extracted = extract_text(document)

    assert extracted.text == "ocr text"
    assert extracted.warnings == ()


def test_extract_text_reads_synthetic_pdf_with_pymupdf(tmp_path: Path) -> None:
    """Verify real PyMuPDF extraction on a synthetic non-sensitive PDF.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.

    Returns
    -------
    None
    """

    import fitz

    person = _person(tmp_path)
    document_dir = person.source_documents
    document_dir.mkdir(parents=True)
    path = document_dir / "20260102 Synthetic Report.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text(
        (72, 72),
        "Synthetic non-sensitive PDF text for PyMuPDF integration testing.",
    )
    pdf.save(path)
    pdf.close()
    document = scan_documents(person)[0]

    extracted = extract_text(document)

    assert "Synthetic non-sensitive PDF text" in extracted.text
    assert extracted.warnings == ()
    assert document_page_count(document) == 1


def test_extract_text_reads_image_only_pdf_with_ocrmypdf(tmp_path: Path) -> None:
    """Verify OCRmyPDF extraction on a synthetic image-only PDF.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.

    Returns
    -------
    None
    """

    if shutil.which("ocrmypdf") is None or shutil.which("tesseract") is None:
        pytest.skip("OCRmyPDF and Tesseract are required for this integration test")

    import fitz

    person = _person(tmp_path)
    document_dir = person.source_documents
    document_dir.mkdir(parents=True)
    source_pdf = tmp_path / "source-text.pdf"
    source = fitz.open()
    page = source.new_page(width=595, height=842)
    page.insert_text(
        (72, 160),
        "SYNTHETIC OCR TEST",
        fontsize=48,
    )
    source.save(source_pdf)
    source.close()

    rendered = fitz.open(source_pdf)
    pixmap = rendered[0].get_pixmap(dpi=250)
    rendered.close()

    image_only_path = document_dir / "20260102 Synthetic OCR Report.pdf"
    image_only = fitz.open()
    image_page = image_only.new_page(width=pixmap.width, height=pixmap.height)
    image_page.insert_image(
        image_page.rect,
        stream=pixmap.tobytes("png"),
    )
    image_only.save(image_only_path)
    image_only.close()
    document = scan_documents(person)[0]

    native_text = documents_module._extract_pdf_text_with_pymupdf(document)
    extracted = extract_text(document)

    assert native_text is not None
    assert not documents_module._has_sufficient_pdf_text(native_text.text)
    assert "SYNTHETIC OCR TEST" in " ".join(extracted.text.split())
    assert extracted.warnings == ()


def test_extract_text_falls_back_to_ocr_when_pymupdf_raises(
    tmp_path: Path,
    monkeypatch,
    capsys,
) -> None:
    """Verify malformed PDFs keep PyMuPDF warning and continue with OCR.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.
    monkeypatch : pytest.MonkeyPatch
        Pytest monkeypatch fixture.
    capsys : pytest.CaptureFixture[str]
        Pytest output capture fixture.

    Returns
    -------
    None
    """

    person = _person(tmp_path)
    document_dir = person.source_documents
    document_dir.mkdir(parents=True)
    path = document_dir / "20260102 Malformed Report.pdf"
    path.write_bytes(b"%PDF-1.7\nnot a valid page tree\n")
    document = scan_documents(person)[0]

    def fake_run(command, **_kwargs):
        sidecar = Path(command[command.index("--sidecar") + 1])
        sidecar.write_text("ocr text after pymupdf failure", encoding="utf-8")
        return subprocess.CompletedProcess(command, 0, "", "")

    monkeypatch.setattr(documents_module.shutil, "which", lambda _: "/usr/bin/ocrmypdf")
    monkeypatch.setattr(documents_module.subprocess, "run", fake_run)

    extracted = extract_text(document)
    captured = capsys.readouterr()

    assert extracted.text == "ocr text after pymupdf failure"
    assert len(extracted.warnings) == 1
    assert extracted.warnings[0].startswith(
        "PyMuPDF non ha potuto estrarre il testo PDF; uso OCRmyPDF se disponibile:"
    )
    assert "MuPDF error" not in captured.out
    assert "MuPDF error" not in captured.err


def test_extract_text_warns_when_pymupdf_raises_without_ocr(
    tmp_path: Path,
    monkeypatch,
    capsys,
) -> None:
    """Verify malformed PDFs are reported when OCR is unavailable.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.
    monkeypatch : pytest.MonkeyPatch
        Pytest monkeypatch fixture.
    capsys : pytest.CaptureFixture[str]
        Pytest output capture fixture.

    Returns
    -------
    None
    """

    person = _person(tmp_path)
    document_dir = person.source_documents
    document_dir.mkdir(parents=True)
    path = document_dir / "20260102 Malformed Report.pdf"
    path.write_bytes(b"%PDF-1.7\nnot a valid page tree\n")
    document = scan_documents(person)[0]

    monkeypatch.setattr(documents_module.shutil, "which", lambda _: None)

    extracted = extract_text(document)
    captured = capsys.readouterr()

    assert extracted.text == ""
    assert len(extracted.warnings) == 1
    assert extracted.warnings[0].startswith(
        "PyMuPDF non ha potuto estrarre il testo PDF; uso OCRmyPDF se disponibile:"
    )
    assert "MuPDF error" not in captured.out
    assert "MuPDF error" not in captured.err


def test_extract_text_keeps_sufficient_pymupdf_text(
    tmp_path: Path,
    monkeypatch,
) -> None:
    """Verify sufficient PyMuPDF text avoids OCR.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.
    monkeypatch : pytest.MonkeyPatch
        Pytest monkeypatch fixture.

    Returns
    -------
    None
    """

    person = _person(tmp_path)
    document_dir = person.source_documents
    document_dir.mkdir(parents=True)
    path = document_dir / "20260102 Report.pdf"
    path.write_bytes(b"%PDF synthetic")
    document = scan_documents(person)[0]

    monkeypatch.setattr(
        documents_module,
        "_extract_pdf_text_with_pymupdf",
        lambda item: documents_module.ExtractedText(
            document_id=item.document_id,
            text="This is enough extracted text from a digital PDF document.",
        ),
    )
    monkeypatch.setattr(documents_module.shutil, "which", lambda _: None)

    extracted = extract_text(document)

    assert extracted.text.startswith("This is enough")
    assert extracted.warnings == ()


def test_extract_text_uses_ocr_when_pymupdf_text_is_insufficient(
    tmp_path: Path,
    monkeypatch,
) -> None:
    """Verify weak PyMuPDF extraction falls back to OCRmyPDF.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.
    monkeypatch : pytest.MonkeyPatch
        Pytest monkeypatch fixture.

    Returns
    -------
    None
    """

    person = _person(tmp_path)
    document_dir = person.source_documents
    document_dir.mkdir(parents=True)
    path = document_dir / "20260102 Report.pdf"
    path.write_bytes(b"%PDF synthetic")
    document = scan_documents(person)[0]

    def fake_run(command, **_kwargs):
        sidecar = Path(command[command.index("--sidecar") + 1])
        sidecar.write_text("ocr text after weak native extraction", encoding="utf-8")
        return subprocess.CompletedProcess(command, 0, "", "")

    monkeypatch.setattr(
        documents_module,
        "_extract_pdf_text_with_pymupdf",
        lambda item: documents_module.ExtractedText(
            document_id=item.document_id,
            text="   \n",
        ),
    )
    monkeypatch.setattr(documents_module.shutil, "which", lambda _: "/usr/bin/ocrmypdf")
    monkeypatch.setattr(documents_module.subprocess, "run", fake_run)

    extracted = extract_text(document)

    assert extracted.text == "ocr text after weak native extraction"
    assert extracted.warnings == ()


def test_extract_text_warns_when_pymupdf_text_is_insufficient_without_ocr(
    tmp_path: Path,
    monkeypatch,
) -> None:
    """Verify weak PyMuPDF extraction warns when OCR is unavailable.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.
    monkeypatch : pytest.MonkeyPatch
        Pytest monkeypatch fixture.

    Returns
    -------
    None
    """

    person = _person(tmp_path)
    document_dir = person.source_documents
    document_dir.mkdir(parents=True)
    path = document_dir / "20260102 Report.pdf"
    path.write_bytes(b"%PDF synthetic")
    document = scan_documents(person)[0]

    monkeypatch.setattr(
        documents_module,
        "_extract_pdf_text_with_pymupdf",
        lambda item: documents_module.ExtractedText(
            document_id=item.document_id,
            text="short",
        ),
    )
    monkeypatch.setattr(documents_module.shutil, "which", lambda _: None)

    extracted = extract_text(document)

    assert extracted.text == "short"
    assert extracted.warnings == (
        "PyMuPDF ha estratto testo insufficiente e non e' disponibile un provider OCR",
    )


def test_extract_text_warns_when_no_pdf_provider_is_available(
    tmp_path: Path,
    monkeypatch,
) -> None:
    """Verify missing PDF providers produce a generic actionable warning.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.
    monkeypatch : pytest.MonkeyPatch
        Pytest monkeypatch fixture.

    Returns
    -------
    None
    """

    person = _person(tmp_path)
    document_dir = person.source_documents
    document_dir.mkdir(parents=True)
    path = document_dir / "20260102 Report.pdf"
    path.write_bytes(b"%PDF synthetic")
    document = scan_documents(person)[0]

    monkeypatch.setattr(
        documents_module, "_extract_pdf_text_with_pymupdf", lambda _: None
    )
    monkeypatch.setattr(documents_module.shutil, "which", lambda _: None)

    extracted = extract_text(document)

    assert extracted.text == ""
    assert extracted.warnings == (
        "Nessun provider di estrazione testo PDF disponibile; installare PyMuPDF "
        "o configurare OCRmyPDF",
    )


def test_extract_text_warns_when_ocrmypdf_fails(
    tmp_path: Path,
    monkeypatch,
) -> None:
    """Verify OCRmyPDF failures are reported without hiding provider availability.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.
    monkeypatch : pytest.MonkeyPatch
        Pytest monkeypatch fixture.

    Returns
    -------
    None
    """

    person = _person(tmp_path)
    document_dir = person.source_documents
    document_dir.mkdir(parents=True)
    path = document_dir / "20260102 Report.pdf"
    path.write_bytes(b"%PDF synthetic")
    document = scan_documents(person)[0]

    def fake_run(command, **_kwargs):
        return subprocess.CompletedProcess(command, 2, "", "ocr failed")

    monkeypatch.setattr(
        documents_module, "_extract_pdf_text_with_pymupdf", lambda _: None
    )
    monkeypatch.setattr(documents_module.shutil, "which", lambda _: "/usr/bin/ocrmypdf")
    monkeypatch.setattr(documents_module.subprocess, "run", fake_run)

    extracted = extract_text(document)

    assert extracted.text == ""
    assert extracted.warnings == (
        "OCRmyPDF non riuscito; estrazione testo PDF saltata: ocr failed",
    )


def test_extract_text_retries_ocrmypdf_without_optimization(
    tmp_path: Path,
    monkeypatch,
) -> None:
    """Verify OCRmyPDF post-processing failures retry without optimization.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.
    monkeypatch : pytest.MonkeyPatch
        Pytest monkeypatch fixture.

    Returns
    -------
    None
    """

    person = _person(tmp_path)
    document_dir = person.source_documents
    document_dir.mkdir(parents=True)
    path = document_dir / "20260102 Report.pdf"
    path.write_bytes(b"%PDF synthetic")
    document = scan_documents(person)[0]
    calls = []

    def fake_run(command, **_kwargs):
        calls.append(command)
        if "--optimize" in command:
            sidecar = Path(command[command.index("--sidecar") + 1])
            sidecar.write_text("retried OCR text", encoding="utf-8")
            return subprocess.CompletedProcess(command, 0, "", "")
        return subprocess.CompletedProcess(
            command,
            2,
            "",
            "Traceback\nOSError: image file is truncated (2 bytes not processed)",
        )

    monkeypatch.setattr(
        documents_module, "_extract_pdf_text_with_pymupdf", lambda _: None
    )
    monkeypatch.setattr(documents_module.shutil, "which", lambda _: "/usr/bin/ocrmypdf")
    monkeypatch.setattr(documents_module.subprocess, "run", fake_run)

    extracted = extract_text(document)

    assert extracted.text == "retried OCR text"
    assert extracted.warnings == ()
    assert calls[0][1:5] == [
        "--skip-text",
        "--output-type",
        "pdf",
        "--continue-on-soft-render-error",
    ]
    assert "--optimize" not in calls[0]
    assert calls[1][1:7] == [
        "--skip-text",
        "--output-type",
        "pdf",
        "--continue-on-soft-render-error",
        "--optimize",
        "0",
    ]


def test_extract_text_summarizes_ocrmypdf_retry_failures(
    tmp_path: Path,
    monkeypatch,
) -> None:
    """Verify OCRmyPDF failures do not dump verbose tool output into warnings.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.
    monkeypatch : pytest.MonkeyPatch
        Pytest monkeypatch fixture.

    Returns
    -------
    None
    """

    person = _person(tmp_path)
    document_dir = person.source_documents
    document_dir.mkdir(parents=True)
    path = document_dir / "20260102 Report.pdf"
    path.write_bytes(b"%PDF synthetic")
    document = scan_documents(person)[0]

    def fake_run(command, **_kwargs):
        if "--optimize" in command:
            return subprocess.CompletedProcess(
                command,
                2,
                "",
                "Page 2\nError: retry failed",
            )
        return subprocess.CompletedProcess(
            command,
            2,
            "",
            "Page 1\nPage 2\nOSError: image file is truncated (4 bytes not processed)",
        )

    monkeypatch.setattr(
        documents_module, "_extract_pdf_text_with_pymupdf", lambda _: None
    )
    monkeypatch.setattr(documents_module.shutil, "which", lambda _: "/usr/bin/ocrmypdf")
    monkeypatch.setattr(documents_module.subprocess, "run", fake_run)

    extracted = extract_text(document)

    assert extracted.text == ""
    assert extracted.warnings == (
        "OCRmyPDF non riuscito; estrazione testo PDF saltata: "
        "OSError: image file is truncated (4 bytes not processed); "
        "retry con --optimize 0 non riuscito: Error: retry failed",
    )


def test_extract_text_locates_ocrmypdf_failure_page_with_bisection(
    tmp_path: Path,
    monkeypatch,
) -> None:
    """Verify OCRmyPDF failure diagnostics bisect source page ranges.

    Parameters
    ----------
    tmp_path : pathlib.Path
        Temporary directory provided by pytest.
    monkeypatch : pytest.MonkeyPatch
        Pytest monkeypatch fixture.

    Returns
    -------
    None
    """

    person = _person(tmp_path)
    document_dir = person.source_documents
    document_dir.mkdir(parents=True)
    path = document_dir / "20260102 Report.pdf"
    path.write_bytes(b"%PDF synthetic")
    document = scan_documents(person)[0]
    diagnostic_ranges = []

    def fake_run(command, **_kwargs):
        if "--pages" not in command:
            if "--optimize" in command:
                return subprocess.CompletedProcess(
                    command,
                    2,
                    "",
                    "Error: retry failed",
                )
            return subprocess.CompletedProcess(
                command,
                2,
                "",
                "OSError: image file is truncated (4 bytes not processed)",
            )
        page_range = command[command.index("--pages") + 1]
        diagnostic_ranges.append(page_range)
        start_text, _, end_text = page_range.partition("-")
        start = int(start_text)
        end = int(end_text or start_text)
        returncode = 2 if start <= 6 <= end else 0
        return subprocess.CompletedProcess(command, returncode, "", "")

    monkeypatch.setattr(
        documents_module, "_extract_pdf_text_with_pymupdf", lambda _: None
    )
    monkeypatch.setattr(documents_module, "_pdf_page_count", lambda _: 8)
    monkeypatch.setattr(documents_module.shutil, "which", lambda _: "/usr/bin/ocrmypdf")
    monkeypatch.setattr(documents_module.subprocess, "run", fake_run)

    extracted = extract_text(document)

    assert extracted.text == ""
    assert extracted.warnings == (
        "OCRmyPDF non riuscito; estrazione testo PDF saltata: "
        "OSError: image file is truncated (4 bytes not processed); "
        "retry con --optimize 0 non riuscito: Error: retry failed; "
        "pagina sorgente non riuscita: 6",
    )
    assert diagnostic_ranges == ["1-4", "5-6", "5", "6"]
